import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def format_doc():
    doc = docx.Document("value_chain_physics_scm_paper_draft_zh_rendered.docx")
    
    # Set Normal style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    rPr = style_normal._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), '宋体')

    # Format all paragraphs and runs
    for p in doc.paragraphs:
        p.paragraph_format.line_spacing = 1.2
        if p.text.startswith("表 1") or p.text.startswith("表 2"):
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.bold = True
                r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')

    # Format tables
    for table in doc.tables:
        table.autofit = True
        for r_idx, row in enumerate(table.rows):
            is_header = (r_idx == 0)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.1
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        if is_header:
                            r.font.bold = True
                            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')

    output_path = "value_chain_physics_scm_paper_draft_zh_rendered.docx"
    try:
        doc.save(output_path)
        print("Successfully saved styled Word document with rendered equations to:", output_path)
    except PermissionError:
        output_path = "value_chain_physics_scm_paper_draft_zh_final.docx"
        doc.save(output_path)
        print("Saved styled Word document with rendered equations to:", output_path)

if __name__ == "__main__":
    format_doc()
