import os
import re
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

base_dir = r"H:\系统科学\价值链物理学"
zh_md = os.path.join(base_dir, "Paper_1_Formal_Proof_OCGS_Qian_Xuesen_ZH.md")
en_md = os.path.join(base_dir, "Paper_1_Formal_Proof_OCGS_Qian_Xuesen_EN.md")

zh_docx = os.path.join(base_dir, "Paper_1_Formal_Proof_OCGS_Qian_Xuesen_ZH.docx")
en_docx = os.path.join(base_dir, "Paper_1_Formal_Proof_OCGS_Qian_Xuesen_EN.docx")

def preprocess_md(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    # Replace file:///.../vcp_axioms_logic_chain.png with relative vcp_axioms_logic_chain.png
    content = re.sub(r'!\[([^\]]*)\]\(file:///[^\)]*vcp_axioms_logic_chain\.png\)', r'![\1](vcp_axioms_logic_chain.png)', content)
    
    # Save temp file
    temp_path = file_path.replace(".md", "_temp.md")
    with open(temp_path, 'w', encoding='utf-8') as f:
        f.write(content)
    return temp_path

def convert_with_pandoc(input_md, output_docx):
    cmd = ["pandoc", input_md, "-o", output_docx, "--from=markdown", "--to=docx"]
    res = subprocess.run(cmd, capture_output=True, text=True, cwd=base_dir)
    print(f"Pandoc output for {output_docx}: {res.returncode}")
    if res.stderr:
        print(f"Pandoc stderr: {res.stderr}")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_docx(docx_path, is_english=False):
    doc = docx.Document(docx_path)
    
    # Set Margins to 1 inch (72 pt)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles styling
    styles = doc.styles
    
    # Normal Paragraph Font
    normal_style = styles['Normal']
    normal_style.font.size = Pt(10.5)
    normal_style.font.name = 'Times New Roman' if is_english else 'Microsoft YaHei'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei' if not is_english else 'SimSun')
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Process Paragraphs
    for p in doc.paragraphs:
        # Check alignment / style
        if p.style.name.startswith('Heading 1'):
            p.style.font.size = Pt(16)
            p.style.font.bold = True
            p.style.font.color.rgb = RGBColor(0x0F, 0x2C, 0x59) # Primary Dark Blue
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif p.style.name.startswith('Heading 2'):
            p.style.font.size = Pt(13)
            p.style.font.bold = True
            p.style.font.color.rgb = RGBColor(0x1B, 0x26, 0x3B)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif p.style.name.startswith('Heading 3'):
            p.style.font.size = Pt(11.5)
            p.style.font.bold = True
            p.style.font.color.rgb = RGBColor(0x41, 0x5A, 0x77)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
        else:
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(4)

    # Process Tables
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Iterate over rows
        for r_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>')) # Prevent row splitting across pages
            
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                if r_idx == 0:
                    set_cell_background(cell, "0F2C59") # Dark Blue Header
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8F9FA") # Light Zebra
                    else:
                        set_cell_background(cell, "FFFFFF")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                            
    doc.save(docx_path)
    print(f"Successfully styled {docx_path}")

def main():
    print("Preprocessing Markdown files...")
    temp_zh = preprocess_md(zh_md)
    temp_en = preprocess_md(en_md)
    
    print("Converting Chinese paper to Word (.docx)...")
    convert_with_pandoc(temp_zh, zh_docx)
    style_docx(zh_docx, is_english=False)
    
    print("Converting English paper to Word (.docx)...")
    convert_with_pandoc(temp_en, en_docx)
    style_docx(en_docx, is_english=True)
    
    # Cleanup temp files
    if os.path.exists(temp_zh):
        os.remove(temp_zh)
    if os.path.exists(temp_en):
        os.remove(temp_en)
        
    print("All Word papers generated and styled successfully!")

if __name__ == "__main__":
    main()
