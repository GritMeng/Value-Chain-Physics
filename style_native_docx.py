import docx, re
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font_eastasia(run, font_name):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), font_name)

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), val)
            border.set(qn('w:sz'), sz)
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), color)
            borders.append(border)
        insideV = OxmlElement('w:insideV')
        insideV.set(qn('w:val'), 'none')
        borders.append(insideV)
        tblPr[0].append(borders)

def process_document(input_filename, output_filename, is_english=True):
    doc = docx.Document(input_filename)
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        
    # Style paragraphs
    for p in doc.paragraphs:
        # Check heading level
        style_name = p.style.name if p.style else ""
        text = p.text.strip()
        
        if style_name.startswith('Heading 1') or (text and text.startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.')) and len(text) < 80):
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = 'Times New Roman' if is_english else 'Times New Roman'
                if not is_english:
                    set_font_eastasia(r, '黑体')
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        elif style_name.startswith('Heading 2') or (text and re.match(r'^\d+\.\d+', text) and len(text) < 80):
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = 'Times New Roman'
                if not is_english:
                    set_font_eastasia(r, '黑体')
                r.font.size = Pt(11.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        else:
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = 'Times New Roman'
                if not is_english:
                    set_font_eastasia(r, '宋体')
                r.font.size = Pt(10.5)

    # Style tables
    for tbl in doc.tables:
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)
        for r_idx, row in enumerate(tbl.rows):
            is_header = (r_idx == 0)
            for c_idx, cell in enumerate(row.cells):
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                if is_header:
                    set_cell_background(cell, "EAECEF")
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.bold = True
                            r.font.size = Pt(9.5)
                else:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9.5)

    doc.save(output_filename)
    omml_count = len(doc._element.xpath('//m:oMath'))
    print(f"Successfully processed {output_filename} with {omml_count} native Word OMML equation objects!")

if __name__ == "__main__":
    import re
    process_document("value_chain_physics_scm_paper_draft_en_native.docx", "value_chain_physics_scm_paper_draft_en.docx", is_english=True)
    process_document("value_chain_physics_scm_paper_draft_zh_native.docx", "value_chain_physics_scm_paper_draft_zh.docx", is_english=False)
